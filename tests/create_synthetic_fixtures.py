#!/usr/bin/env python3
"""Create privacy-safe DOCX and XLSX fixtures for local smoke tests."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from openpyxl import Workbook


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_heading("Synthetic verification", 1)
    document.add_paragraph("No personal data.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "예시"
    table.cell(1, 1).text = "1"
    document.save(output / "sample.docx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "양식"
    sheet["A1"] = "항목"
    sheet["B1"] = "값"
    sheet["A2"] = "예시"
    sheet["B2"] = "=1+1"
    sheet.merge_cells("A3:B3")
    sheet.print_area = "A1:B3"
    hidden = workbook.create_sheet("숨김")
    hidden.sheet_state = "hidden"
    workbook.save(output / "sample.xlsx")
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
